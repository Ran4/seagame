from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path("reports/Sea_Game_Review_2026-03-27.docx")


@dataclass
class Finding:
    severity: str
    title: str
    summary: str
    evidence: list[str]
    impact: str
    recommendation: str


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(11)

    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(20, 54, 96)

    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(44, 82, 130)

    styles["Heading 3"].font.name = "Aptos"
    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Sea Game Review - page ")
    add_page_number(footer)


def add_title_page(doc: Document) -> list[str]:
    words: list[str] = []

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_after = Pt(18)
    run = p.add_run("Sea Game Review")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(20, 54, 96)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Code and Gameplay Assessment")
    run.font.name = "Aptos Display"
    run.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(18)
    p.space_after = Pt(18)
    run = p.add_run("Repository: /home/ran/seagame\nReview date: 2026-03-27")
    run.font.size = Pt(11)

    intro = dedent(
        """
        Scope: review of the current 2D gameplay, simulation, content systems, and production-readiness of the codebase. The in-progress 3D renderer was intentionally excluded except where it leaks into shared runtime behavior such as bootstrapping, build stability, or input routing. The review is based on direct inspection of 48 TypeScript source files totaling roughly 10,869 lines, the implemented feature notes, the initial design document, and a production build check via npm run build.
        """
    ).strip()
    doc.add_paragraph(intro)
    words.append(intro)

    bullets = [
        "Primary goal of the report: identify the highest-value fixes and the most important gameplay adjustments before more content is layered on top.",
        "Review method: static analysis of the game loop, crew AI, harbor systems, command system, world map, notices, items, and supporting feature documents.",
        "Notable constraint: 3D rendering work is still in progress, so the review does not score the visual renderer itself.",
    ]
    for bullet in bullets:
        doc.add_paragraph(bullet, style="List Bullet")
        words.append(bullet)

    doc.add_page_break()
    return words


def add_heading(doc: Document, text: str, level: int) -> str:
    doc.add_heading(text, level=level)
    return text


def add_paragraphs(doc: Document, paragraphs: list[str], words: list[str]) -> None:
    for paragraph in paragraphs:
      clean = dedent(paragraph).strip()
      if not clean:
          continue
      doc.add_paragraph(clean)
      words.append(clean)


def add_bullets(doc: Document, items: list[str], words: list[str]) -> None:
    for item in items:
        clean = dedent(item).strip()
        doc.add_paragraph(clean, style="List Bullet")
        words.append(clean)


def add_findings_table(doc: Document, findings: list[Finding]) -> list[str]:
    words: list[str] = []
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "Severity", True)
    set_cell_text(hdr[1], "Finding", True)
    set_cell_text(hdr[2], "Why it matters", True)
    set_cell_text(hdr[3], "Key evidence", True)

    for finding in findings:
        row = table.add_row().cells
        set_cell_text(row[0], finding.severity)
        set_cell_text(row[1], finding.title)
        set_cell_text(row[2], finding.summary)
        set_cell_text(row[3], "; ".join(finding.evidence))
        words.extend([finding.severity, finding.title, finding.summary, " ".join(finding.evidence)])

    return words


def build_findings() -> list[Finding]:
    return [
        Finding(
            severity="High",
            title="Production build is currently broken",
            summary=(
                "The checked-in game does not complete a production build. "
                "This blocks deployment, packaged playtests, and any confidence that gameplay changes are shipping in a stable form."
            ),
            evidence=[
                "Build check: npm run build",
                "src/main.ts:14",
                "Build failure: top-level await is not available for current Vite/esbuild target",
            ],
            impact=(
                "This is more than a tooling nuisance. It removes the fastest feedback loop for every other system in the project. "
                "When the build is red, it becomes harder to separate real gameplay regressions from local development quirks, and it raises the cost of integrating new features."
            ),
            recommendation=(
                "Move startup into an async bootstrap function or raise the build target deliberately. "
                "The safer option is to wrap loadConfig in a bootstrap function and keep the browser target unchanged."
            ),
        ),
        Finding(
            severity="High",
            title="Harbor NPCs dilute the mutiny denominator",
            summary=(
                "Mutiny threshold is based on all human actors, but harbor NPCs are also human actors. "
                "That means docked play quietly raises the number of mutineers required for a mutiny, sometimes beyond what the actual crew can even reach."
            ),
            evidence=[
                "src/crew/update.ts:800-809",
                "src/harbor.ts:427-453",
            ],
            impact=(
                "This weakens one of the game’s only true fail states exactly when the harbor layer is supposed to be a strategic release valve. "
                "For example, with four real crew and four human NPCs in town, the 60 percent threshold becomes five mutineers, which the crew can never produce."
            ),
            recommendation=(
                "Filter mutiny logic to controllable, non-NPC human crew only. "
                "The same rule should be reused anywhere crew-wide thresholds are calculated."
            ),
        ),
        Finding(
            severity="High",
            title="Recruit sailor is not really limited per visit; it is effectively limited for the whole run",
            summary=(
                "The innkeeper menu checks whether any actor has the status recruited_this_visit, but that status is never cleared on undock or on a new harbor visit."
            ),
            evidence=[
                "src/game.ts:95",
                "src/menu.ts:125-128",
                "Search result: recruited_this_visit appears only in those two places",
            ],
            impact=(
                "The player-facing promise is one recruit per visit. The implemented behavior is closer to one recruit total, or one recruit until that actor leaves the active actor list. "
                "That will read as a bug rather than a design choice."
            ),
            recommendation=(
                "Clear the visit flag on undock, or store the limitation in docking state keyed to island visit rather than on the recruited actor."
            ),
        ),
        Finding(
            severity="Medium",
            title="Hidden-island discovery state is mostly dead data",
            summary=(
                "The game tracks spottedIslands and updates it from lookouts and notices, but the map overlay and map click handler do not consult that set when deciding whether a hidden island is visible or targetable."
            ),
            evidence=[
                "src/notices.ts:52-64",
                "src/crew/update.ts:826-840",
                "src/render/map.ts:58-63",
                "src/worldmap.ts:134",
            ],
            impact=(
                "Two discovery systems currently fail to create actionable discovery. "
                "A lookout can cry Land ho and the notice board can mark a hidden island, but unless an expert navigator is already present the newly discovered island still stays hidden on the actual map."
            ),
            recommendation=(
                "Treat spottedIslands as the canonical visibility set for hidden islands, and let expert navigation widen or accelerate discovery rather than bypass the set entirely."
            ),
        ),
        Finding(
            severity="Medium",
            title="The notice board claims truthfulness while advertising systems that do not exist",
            summary=(
                "The notice board text asserts that it only references real mechanics, but several notices imply restocking, shipwright services, and other harbor affordances that are not implemented in the current game."
            ),
            evidence=[
                "src/notices.ts:3-21",
                "src/game.ts:555-561",
                "features/implemented/2026-03-23_HARBOR_TOWNS.md",
            ],
            impact=(
                "This is a content integrity problem. "
                "False affordances are especially damaging in simulation games because players spend time testing what the world will and will not respond to."
            ),
            recommendation=(
                "Either cut those lines for now or connect them to placeholder but real mechanics such as a one-click resupply, repair, or shop transaction."
            ),
        ),
        Finding(
            severity="Medium",
            title="Need satisfaction is free, so the survival layer does not create resource pressure",
            summary=(
                "Eating and sleeping fully restore key needs without consuming ship resources, reserving facilities, or creating opportunity cost beyond time. "
                "As a result, hunger and fatigue act as animation routing rather than as a management system."
            ),
            evidence=[
                "src/crew/update.ts:346-363",
                "src/items.ts:27-46",
            ],
            impact=(
                "The game looks like a ship management sim, but the current rules behave more like a self-healing toy box. "
                "That undermines pacing, reduces meaningful planning, and weakens morale because morale inputs are not backed by logistics."
            ),
            recommendation=(
                "Introduce at least one concrete supply loop soon: meal stock, water, rum, or galley throughput. "
                "The smallest viable version would be limited food items consumed on Eat, with morale and mutiny reacting to shortages."
            ),
        ),
        Finding(
            severity="Medium",
            title="Grog is currently a nearly free answer to morale problems",
            summary=(
                "A grog ration can be granted immediately from the bartender, drawn for free from harbor barrels, and then converted directly into a morale bump. "
                "Without money, stock scarcity, or stronger drawbacks, the system collapses a lot of tension."
            ),
            evidence=[
                "src/game.ts:540-546",
                "src/crew/update.ts:552-560",
                "src/crew/update.ts:977-990",
            ],
            impact=(
                "This flattens the central risk curve. "
                "If morale is supposed to be the umbrella system tying the game together, its strongest quick fix cannot also be nearly consequence-free."
            ),
            recommendation=(
                "Put grog behind either cost, scarce stock, or stronger knock-on risks such as worse obedience, accidents, fights, or reduced work output."
            ),
        ),
    ]


def write_document() -> int:
    doc = Document()
    configure_document(doc)

    findings = build_findings()
    words: list[str] = []

    words.extend(add_title_page(doc))

    add_heading(doc, "Executive Summary", 1)
    add_paragraphs(
        doc,
        [
            """
            Sea Game already has the part that matters most for this genre: a recognizable identity. The strongest aspect of the project is not any single mechanic, but the way the current systems already produce a specific fantasy. A small pirate crew wanders a ship, tends to itself, sings shanties, gets drunk, pets animals, visits harbor towns, and occasionally destabilizes into melodrama. That is a much better starting point than a technically tidy project with no point of view. The ship feels inhabited, the harbor update clearly improved the social density of the world, and the activity log plus bubble-based feedback do a lot of work in selling the simulation.
            """,
            """
            The project’s biggest weakness is that the simulation currently overpromises consequence. A lot of visible systems imply management depth, but many of them resolve for free. Crew eat without finite food. Crew sleep without any real berth pressure. Grog is handed out with little cost. Harbor content offers flavor and mood, but not yet enough economic or strategic tradeoffs. Because of that, the game reads well moment to moment but does not yet accumulate much long-form pressure. The result is a sandbox that is charming in five-minute windows but currently thin in thirty-minute windows.
            """,
            """
            On the code side, the project is in better shape than a purely experimental jam prototype, but it is also entering the phase where central state ownership needs to become more deliberate. Most gameplay logic still lives in a few very large files with broad authority, especially src/game.ts and src/crew/update.ts. That is workable for now, but only while the behavior set remains legible in one person’s head. Several of the highest-value defects come from cross-system assumptions that are not encoded centrally: who counts as crew, when a visit-scoped status should reset, and what a discovered island actually means to the rest of the game.
            """,
            """
            My overall judgment is that the game has a strong fantasy layer, a promising autonomous-crew foundation, and enough implemented personality to justify further investment. The next step should not be breadth for its own sake. The next step should be to tighten consequence and clarity around the existing loop. In practical terms that means fixing the production build, making crew-only rules actually crew-only, making discovery actionable, and turning at least one need system into a real resource-management problem before adding a new major content pillar.
            """,
        ],
        words,
    )

    add_heading(doc, "Highest-Priority Findings", 1)
    words.extend(add_findings_table(doc, findings))

    for finding in findings:
        add_heading(doc, f"{finding.severity}: {finding.title}", 2)
        add_paragraphs(
            doc,
            [
                finding.summary,
                f"Evidence: {'; '.join(finding.evidence)}.",
                finding.impact,
                f"Recommended action: {finding.recommendation}",
            ],
            words,
        )

    doc.add_page_break()

    add_heading(doc, "What Is Working Well", 1)
    add_paragraphs(
        doc,
        [
            """
            The game’s clearest success is that it already produces stories without needing authored quest content. A crew member can get tired, wander to a bed, miss a conversation, drink grog, join a dance, argue, and then help trigger a morale swing. That kind of emergent sequence is the right foundation for the project’s stated ambition of being Dwarf Fortress at sea. It means the core attraction is not dependent on cinematics or branching dialogue trees. It comes from systems touching each other in visible ways.
            """,
            """
            The harbor update is also directionally correct. It gives the player a change of texture, better social density, and more chances to read the ship as part of a larger world rather than as a sealed playset. The tavern, inn, market, smithy, townsfolk, and harbor cat are modest additions in strict mechanical terms, but they improve the game’s readability and mood disproportionately. This is important because management games live or die by whether routine activity stays pleasant to watch.
            """,
            """
            Another strength is the project’s use of low-cost feedback channels. The activity log is simple, but it makes state changes legible. Speech bubbles, thought bubbles, and short status labels do not just add charm; they compress simulation state into something the player can parse quickly. The project also already understands reward cadence at a small scale. Shanties, dog morale bonuses, gossip after NPC conversations, and harbor arrival celebrations all function as tiny emotional payouts that keep autonomous systems from feeling cold.
            """,
            """
            Finally, the feature documentation is unusually helpful for a project at this stage. The implemented feature notes explain intent and constraints, and the issues folder shows where the design already knows it is incomplete. That is valuable because it makes it easier to tell the difference between a missing system, a temporary placeholder, and a bug that contradicts the intended design.
            """,
        ],
        words,
    )

    add_heading(doc, "Codebase Assessment", 1)
    add_heading(doc, "State Ownership and File Structure", 2)
    add_paragraphs(
        doc,
        [
            """
            The codebase is still in the productive early-mid phase where large files are not automatically a problem, but the warning signs are visible. src/game.ts owns bootstrapping, UI state transitions, docking flow, input dispatch, context-menu execution, order polling, spoilage ticking, and actor update orchestration. src/crew/update.ts owns needs decay, morale, drunkenness, lust, state transitions, shanties, dancing, tavern brawls, mutiny, and multiple layers of autonomous behavior. Neither file is disorganized in a chaotic sense, but both have become high-gravity files where new features naturally accumulate.
            """,
            """
            The practical risk is not just size. It is hidden coupling. When visit-scoped logic, actor classification, map discovery, and mood systems all live in adjacent but separate blocks, it becomes easy to add a feature that works locally while violating a global assumption. The recruited_this_visit bug is a good example. Nothing in that implementation is syntactically wrong. The problem is that the lifecycle of the flag is not owned anywhere. The same pattern appears in mutiny, where the game has a concept of crew in some places and a concept of human actors in others.
            """,
            """
            I would not recommend a grand architecture rewrite yet. The better move is to start extracting ownership boundaries around the systems that already produce defects. A crew-classification helper, a discovery/visibility helper for islands, and a visit-scoped harbor state helper would remove multiple classes of bugs immediately. The goal is to make the rules harder to express incorrectly, not simply to move functions into smaller files for aesthetic reasons.
            """,
        ],
        words,
    )

    add_heading(doc, "Simulation and AI Logic", 2)
    add_paragraphs(
        doc,
        [
            """
            The crew AI is the heart of the project, and it is already good enough to justify the rest of the game. The idle update chain has a readable priority order: satisfy needs, seek lust partners, manage lanterns, drink, socialize, perform celebratory group actions, then wander. That is a solid pattern for a life-sim-style state machine because it produces recognizable routines without making every actor look deterministic.
            """,
            """
            The weak point is not behavior variety. It is constraint depth. The AI can decide to eat, but eating is just a time sink attached to a stove tile. The AI can decide to sleep, but beds do not form a scarce resource unless multiple actors collide incidentally. The AI can decide to take grog, but rum is effectively abundant. This means the decision tree is active, but the economy under the decision tree is still thin. The result is a lively simulation whose outcomes often do not matter enough.
            """,
            """
            There is also a subtle readability issue in how many systems silently modify morale. Morale is influenced by needs, friendship, dogs, injuries, harbor bonuses, darkness, grog, dancing, singing, island spotting, NPC conversations, and tavern brawls. That is not inherently bad. In fact it is close to what the game wants. The problem is that the player currently gets only fragments of explanation. When orders fail because morale is low, the simulation will feel fair only if the player can quickly reconstruct why morale got low and which levers are real fixes instead of cosmetic ones.
            """,
        ],
        words,
    )

    add_heading(doc, "Content Integrity", 2)
    add_paragraphs(
        doc,
        [
            """
            The project benefits from having a strong voice, but content credibility matters in systemic games. The notice board is the clearest example. Its comment states that notices reference only things that actually exist in the game, yet some notices still imply mechanics that the current build does not support. A player may tolerate placeholder text in an action game, but in a simulation game those lines function like promises. If the world says shipwrights or restocking are present, players will spend attention trying to use them.
            """,
            """
            The same issue applies to hidden-island discovery. A lookout shouting Land ho and a notice board marking a hidden island should create a new tactical option. Right now those actions mostly create a line in the log and a morale spike. Mechanically, the world has not actually opened up in a usable way. That weakens one of the most satisfying kinds of simulation payoff: knowledge turning into capability.
            """,
        ],
        words,
    )

    add_heading(doc, "Gameplay Review", 1)
    add_heading(doc, "Fantasy and Player Appeal", 2)
    add_paragraphs(
        doc,
        [
            """
            The game’s fantasy is already legible within seconds, and that is a real achievement. You are not manually piloting a ship in an arcade sense. You are the invisible coordinator of a pirate micro-society. That framing gives the game room to be funny, dirty, affectionate, and slightly chaotic without losing coherence. The dogs, parrot, monkey, tavern brawls, shanties, gossip, and harbor visits all support that fantasy.
            """,
            """
            The strongest emotional promise of the game is not conquest. It is attachment. Crew members have names, evolving relationships, changing morale, and accumulating skills. That is exactly the right axis for this kind of simulation. The more the game can make a specific sailor matter mechanically and socially, the more every good or bad event lands. The codebase is already pushing in this direction through skills, traits, relation values, and death handling.
            """,
        ],
        words,
    )

    add_heading(doc, "The Current Core Loop", 2)
    add_paragraphs(
        doc,
        [
            """
            In its present form, the core loop is best described as observational stewardship. The player watches the crew self-regulate, steps in with right-click orders when necessary, occasionally routes navigation, and uses harbor visits to refresh mood and texture. That loop is enjoyable in short bursts because the ship always has some amount of ambient life. There is usually something to notice, whether it is a conversation, a need-driven behavior, or a small dramatic event.
            """,
            """
            The weakness is that the loop does not yet escalate. The player can make local interventions, but there are not enough medium-term commitments. There is no meaningful resource plan for the next day, no economic reason to prefer one harbor over another beyond flavor, and no strong strategic distinction between a well-run ship and a merely surviving one. Progress exists mostly as incremental skill growth and more actors in the simulation, not as a changing problem space.
            """,
            """
            This is why the current game feels better as a toy than as a campaign. The toy is good. The campaign scaffolding is not ready yet. That is not a criticism of ambition; it is a prioritization point. Before adding storms, combat, treasure maps, and broader world systems, the project needs one or two existing loops to become truly binding on player decisions.
            """,
        ],
        words,
    )

    add_heading(doc, "Shipboard Play", 2)
    add_paragraphs(
        doc,
        [
            """
            Shipboard life is the most successful part of the game. The layout is readable, the deck switching logic matches the fantasy, and the crew’s autonomous wandering gives the ship a sense of habitation. Small flavor systems such as lantern lighting, animal interactions, and social chatter work disproportionately well because they turn dead space into observed routine.
            """,
            """
            Where shipboard play underdelivers is in hard choices. Crew can eat forever because food is not actually consumed. They can sleep without berth ownership or work scheduling. A player may enjoy watching the galley and bunks get used, but the simulation is not yet forcing prioritization. This matters because the ship is supposed to be where roughly eighty percent of play happens. If that eighty percent is mostly decorative self-maintenance, the game risks becoming passive despite having many moving parts.
            """,
            """
            The most important improvement here would be to convert one shipboard routine into a constrained system. Food is the cleanest candidate. Even a simple stock number, coupled with barrel contents and spoilage, would immediately make route choice, harbor resupply, and crew well-being feel connected. That single change would make the existing hunger, morale, and harbor systems all pull in the same direction.
            """,
        ],
        words,
    )

    add_heading(doc, "Harbor Play", 2)
    add_paragraphs(
        doc,
        [
            """
            Harbors are currently more content-rich than the sea layer, and that is both a strength and a warning sign. They provide a strong visual and social reset. The four-building layout is compact but legible, NPC roles are immediately understandable, and the harbor cat is exactly the sort of small touch that makes a place memorable.
            """,
            """
            Mechanically, however, the harbor is still closer to a morale spa than to an economy hub. Buy grog has no visible price. Browse wares is flavor text. Recruitment is useful but bugged in scope. Smithy functionality is implied but absent. Rest at the inn is atmospheric, yet because sleep is free everywhere it is not meaningfully different from shipboard sleep except as scenery. The result is that harbors feel good to enter but do not yet create enough strategic contrast.
            """,
            """
            The fix is not to add ten more harbor interactions. It is to make two or three existing ones materially consequential. Recruitment should cost something and create new social/skill possibilities. Merchant interactions should alter supply state. One harbor service should directly relieve a shipboard pressure that the player genuinely feels during voyages. Once those links exist, the current town layout will punch far above its content budget.
            """,
        ],
        words,
    )

    add_heading(doc, "World Map and Sailing", 2)
    add_paragraphs(
        doc,
        [
            """
            The sailing layer is functionally competent but emotionally thin. Navigation and helmsman separation is a nice simulation touch, and the map overlay is clear. The problem is that the sea between islands has too little decision density at the moment. Outside of discovery pings and the anticipation of docking, long stretches of sailing do not present enough changing risk or reward.
            """,
            """
            This makes harbor content carry too much of the game’s variety burden. The player should feel some tension while away from port: dwindling supplies, mood drift, route tradeoffs, weather risk, or crew specialization mattering to the journey itself. Without that, selecting a destination can become a mostly aesthetic choice instead of a strategic one.
            """,
            """
            The hidden-island system is especially important here because discovery is one of the easiest ways to make travel intrinsically exciting. At the moment the project has the skeleton of discovery but not the loop closure. Fixing spottedIslands so that discovery changes map visibility would improve both the lookout role and the sailing layer without needing a large content drop.
            """,
        ],
        words,
    )

    add_heading(doc, "Agency, Clarity, and Friction", 2)
    add_paragraphs(
        doc,
        [
            """
            The game generally succeeds at being readable, but player agency is still narrower than the feature surface suggests. The player can issue commands, yet many core needs resolve automatically and many world interactions are one-step actions. Because of that, the player often feels like a spectator with interruption privileges rather than a captain shaping a ship culture.
            """,
            """
            There is nothing wrong with low-micromanagement by itself. In fact it is part of the appeal here. The missing ingredient is higher-level control. The game wants role assignment, preferred jobs, or priority settings more than it wants more click-to-do-X interactions. Telling the system what matters should become more important than manually telling a sailor where to stand.
            """,
            """
            The clearest friction point today is explanation. A sailor can refuse orders due to low morale, but the game does not yet expose a concise rationale chain for why that sailor is currently unhappy and what action will best help. For a simulation-heavy design, explanation is part of agency. If players cannot form a reliable mental model, the game becomes charming but opaque.
            """,
        ],
        words,
    )

    add_heading(doc, "Pacing and Progression", 2)
    add_paragraphs(
        doc,
        [
            """
            Progression exists in fragments: skills rise, new recruits join, relations drift, hidden islands can theoretically be discovered, and mood swings create local drama. What is missing is a stronger sense of compounding ship identity. A good progression system for this game would make the ship feel increasingly specific over time: particular specialists, recurring shortages, favorite routines, and a growing web of consequences from prior choices.
            """,
            """
            Right now the biggest pacing risk is plateau. After the player understands how to click around the ship, visit a harbor, and watch crew behavior, there is not yet a strong second layer of mastery to climb. That is where supplies, repairs, route efficiency, reputation, or role assignment could do real work. The project does not need all of those. It needs one of them implemented deeply enough that the rest of the systems begin orbiting it.
            """,
        ],
        words,
    )

    add_heading(doc, "Tone and Audience Fit", 2)
    add_paragraphs(
        doc,
        [
            """
            The game has a bold tone rather than a generic one, and that is usually an asset. The mix of pirate melodrama, autonomous life sim behavior, animals, gossip, drinking, and sexual absurdity makes the project memorable. It does not read like an interchangeable colony sim reskin.
            """,
            """
            That said, the sexual-content layer is currently more mechanically developed than some of the more universally legible management layers. Barrel copulation, semen items, and lust systems are not inherently a problem, but they do shape first impressions and audience expectations. If the long-term goal is a broad management-sim audience, the project should be careful that its strongest management loops do not feel secondary to its most provocative jokes.
            """,
            """
            The right answer here is not necessarily to remove that content. It is to make sure the ship-management fantasy remains the primary pillar and that the more explicit material stays additive rather than defining. Right now the balance is close enough that it is worth watching deliberately.
            """,
        ],
        words,
    )

    add_heading(doc, "Recommended Roadmap", 1)
    add_heading(doc, "Immediate Fixes Before New Feature Breadth", 2)
    add_bullets(
        doc,
        [
            "Repair the production build by removing the unsupported top-level await boot path in src/main.ts.",
            "Filter mutiny, crew counts, and any other crew-wide thresholds to non-NPC controllable humans only.",
            "Fix recruit_sailor scoping so the one-per-visit rule resets correctly.",
            "Make spottedIslands actually drive map visibility and destination selection for hidden islands.",
            "Trim or rewrite notice-board text so every claimed affordance is genuinely usable in the current build.",
        ],
        words,
    )

    add_heading(doc, "Highest-Leverage Gameplay Upgrade", 2)
    add_paragraphs(
        doc,
        [
            """
            If I had to choose only one deeper gameplay upgrade for the next milestone, it would be a real food-and-supplies loop. That single addition would strengthen shipboard management, justify harbor trade, give the player a reason to care about storage and spoilage, make morale more legible, and create route pressure on the world map. It is the smallest change that touches the most existing systems without requiring combat or a major content explosion.
            """,
            """
            A minimal version would look like this: barrels hold finite provisions, Eat consumes one, harbor markets can restock them, low stock applies morale pressure, and the UI exposes remaining days of food. That is enough to transform several decorative systems into operational ones.
            """,
        ],
        words,
    )

    add_heading(doc, "Second-Layer Agency Upgrade", 2)
    add_paragraphs(
        doc,
        [
            """
            After supply pressure, the next best investment is higher-level crew control. The project wants job priorities, preferred roles, or standing orders more than it wants more one-off context-menu verbs. This would preserve the low-micromanagement identity while increasing strategic authorship. Players should be deciding who is primarily a helmsman, who is a cook, who should avoid rum, who gets first claim on sleep, and who should prioritize lantern duty.
            """,
            """
            That kind of system would also make the existing skill framework pay off more cleanly. Skills are already present, but the player cannot yet express enough long-term preference around them. A role or priority layer would turn skills from passive metadata into planning tools.
            """,
        ],
        words,
    )

    add_heading(doc, "Content Additions That Would Land Best After the Above", 2)
    add_bullets(
        doc,
        [
            "Storms, because they would become meaningful once supplies, morale, and specialized crew roles already matter.",
            "Trading, because it would have obvious interaction with supply scarcity and harbor differentiation.",
            "Ship-to-ship combat, because the game would then have clearer stakes for crew loss, injuries, repairs, and reputation.",
            "Per-island harbor variants, because stronger mechanics would let content variation express strategy rather than only flavor.",
        ],
        words,
    )

    add_heading(doc, "Conclusion", 1)
    add_paragraphs(
        doc,
        [
            """
            Sea Game is already interesting. That is the hardest milestone. The code and the gameplay both show a project with a strong internal voice, a good autonomous-simulation core, and multiple small systems that already generate memorable moments. The current risk is not that the game is dull. The current risk is that new breadth arrives before the existing loops are made binding enough to support it.
            """,
            """
            The fastest path to a meaningfully better game is to tighten consequences around what already exists. Fix the build. Fix crew classification. Make discovery actionable. Stop promising mechanics the world cannot answer. Then choose one shipboard resource loop and make it real. If those steps happen first, the next wave of features will have something solid to attach to. If they do not, the project may keep getting more colorful while remaining structurally light.
            """,
            """
            My overall recommendation is therefore simple: consolidate the current fantasy into a stronger management game before expanding the fantasy outward. The project has earned that discipline.
            """,
        ],
        words,
    )

    doc.add_page_break()

    add_heading(doc, "Appendix A: Reviewed Materials", 1)
    add_bullets(
        doc,
        [
            "Initial design document: architecture/initial_GDD.md",
            "Gameplay core: src/game.ts, src/crew/update.ts, src/crew/commands.ts, src/crew/movement.ts",
            "World and docking: src/worldmap.ts, src/harbor.ts, src/ship.ts",
            "Content layers: src/conversation.ts, src/notices.ts, src/items.ts",
            "Supporting docs: features/implemented/*.md, issues/*.md",
            "Build validation: npm run build",
        ],
        words,
    )

    add_heading(doc, "Appendix B: Most Important Evidence References", 1)
    evidence_refs = [
        "Build failure source: src/main.ts:14",
        "Mutiny denominator: src/crew/update.ts:800-809",
        "NPC actor creation: src/harbor.ts:427-453",
        "Recruit flag write: src/game.ts:95",
        "Recruit flag read: src/menu.ts:125-128",
        "Notice-based island reveal: src/notices.ts:52-64",
        "Map overlay hidden-island gate: src/render/map.ts:61",
        "Map click hidden-island gate: src/worldmap.ts:134",
        "Free eating restore: src/crew/update.ts:346-352",
        "Free sleep restore: src/crew/update.ts:357-363",
        "Free bartender grog grant: src/game.ts:540-546",
        "Harbor grog scavenging: src/crew/update.ts:977-990",
    ]
    add_bullets(doc, evidence_refs, words)

    add_heading(doc, "Appendix C: System-by-System Notes", 1)
    add_heading(doc, "Crew Identity and Attachment", 2)
    add_paragraphs(
        doc,
        [
            """
            The project is doing the correct thing by making crew members mechanically distinct before it chases large-scale content volume. Skills, relationships, sex, traits, injuries, drunkenness, lust, and animal preferences all push the crew toward being memorable rather than fungible. That is exactly the right design instinct. In management sims, the player does not form attachment because the writer says a character is important. The player forms attachment because a character occupies a useful and visible place inside the ship’s routine. Sea Game is already close to that. Jack with a cutlass, Mary with grog in inventory, a dog-loved sailor, a future expert navigator, and a singer who keeps morale high are all the beginnings of a roster the player can care about.
            """,
            """
            The next improvement here should be stronger expression, not more raw variables. The game already has enough hidden state to create personality; it now needs better ways for the player to read and act on that personality. Small UI touches would go a long way: show each sailor’s strongest skill, show top friends and rivals, and show the biggest current contributors to morale. None of that requires a new subsystem. It just turns existing data into decisions.
            """,
        ],
        words,
    )

    add_heading(doc, "Commands and Control Surface", 2)
    add_paragraphs(
        doc,
        [
            """
            The command surface is flexible enough for the current game. Context menus, self-actions, actor-to-actor interactions, deck transfers, and typed commands create a good amount of expressive control. The strength of this approach is that it keeps the ship feeling tactile. The player points at the world and issues contextual intent rather than manipulating abstract task cards.
            """,
            """
            The weakness is scalability. As more systems arrive, purely contextual actions can become noisy unless they are supported by higher-order controls. The game already hints at this through the existence of command queues and order polling. That suggests the design wants to grow into scripting, delegation, and indirect control. I think that is correct. The next step should be to let the player set persistent priorities or roles so that the command surface handles exceptions while standing orders handle routine life aboard the ship.
            """,
        ],
        words,
    )

    add_heading(doc, "Economy and Resource Shape", 2)
    add_paragraphs(
        doc,
        [
            """
            The current resource model is the game’s most important missing middle layer. There are already barrels, inventories, spoilage hooks, harbor merchants, grog, and need bars. Those ingredients are enough to support an actual economy, but they are not yet wired into one another tightly enough. This is promising because it means the project does not need to invent a brand-new pillar. It needs to finish connecting the one it already started.
            """,
            """
            If the team wants to stay pragmatic, it should resist the temptation to add gold, ten trade goods, and dynamic market prices all at once. The first version should be intentionally small. Food, grog, and maybe one repair resource are enough. What matters is that consumption, scarcity, acquisition, and morale all connect. Once that loop is stable, additional cargo types and harbor specialization will feel like expansion rather than patching holes.
            """,
        ],
        words,
    )

    add_heading(doc, "Testing and Release Discipline", 2)
    add_paragraphs(
        doc,
        [
            """
            The production build failure is a symptom of a broader issue: gameplay code is moving quickly, but there is not yet a clear release gate around the shared runtime. This is normal for an early game prototype, but the project is now large enough that a small amount of discipline will pay for itself. At minimum, the build should be green at the end of each feature slice, and a small smoke-check list should exist for the major loops: boot game, select crew, issue command, sail, dock, undock, talk to NPC, recruit, and trigger notice board.
            """,
            """
            I am not recommending heavy automated coverage across the whole game right now. That would probably slow iteration too much relative to the current design churn. I am recommending targeted validation around the systems that are already expensive to debug manually: startup, docking state transitions, actor classification, and map visibility. A few focused tests or assertions would catch exactly the kind of bugs identified in this report.
            """,
        ],
        words,
    )

    add_heading(doc, "Where New Content Will Help Most", 2)
    add_paragraphs(
        doc,
        [
            """
            The best future content for this game is content that sharpens existing choices rather than bypassing them. Storms would be excellent because they pressure navigation, morale, supplies, lighting, and role quality all at once. Repairs would be excellent because they give harbors and injuries more meaning. Trading would be excellent because it forces route planning and turns island flavor into economic differentiation.
            """,
            """
            By contrast, content that only adds more ambient scenes should be treated as secondary until the management spine is firmer. More conversation lines, more decorative harbor variants, and more one-off jokes will still be welcome later, but they will land better once the ship’s day-to-day operation produces stronger consequences. The existing simulation already has enough flavor to survive a period of mechanical deepening.
            """,
        ],
        words,
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)

    return sum(len(text.split()) for text in words)


if __name__ == "__main__":
    word_count = write_document()
    print(f"Wrote {OUTPUT_PATH}")
    print(f"Approximate word count: {word_count}")
